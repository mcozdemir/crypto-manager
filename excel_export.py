# -*- coding: utf-8 -*-
"""
excel_export.py
----------------
Tarama sonuçlarını biçimlendirilmiş bir .xlsx dosyasına aktarır.
3 sekme oluşturur: "Tüm Sonuçlar", "En Güçlü LONG", "En Güçlü SHORT".

Not: Bu bir veri dökümü (statik değerler) olduğundan formül kullanılmaz;
sayılar tarama anındaki hesaplanmış sonuçlardır ve tekrar tarama
yapılmadan değişmez.
"""

import io
from datetime import datetime
from typing import List, Dict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
LONG_FILL = PatternFill(start_color="E6F7F5", end_color="E6F7F5", fill_type="solid")
SHORT_FILL = PatternFill(start_color="FDECEC", end_color="FDECEC", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"), right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"), bottom=Side(style="thin", color="D9D9D9"),
)

COLUMNS = [
    ("symbol", "Coin", 12),
    ("timeframe", "Zaman Dilimi", 12),
    ("pattern", "Formasyon", 26),
    ("direction", "Yön", 8),
    ("score", "Skor", 8),
    ("success_probability", "Başarı %", 10),
    ("breakout", "Breakout", 10),
    ("volume_desc", "Hacim", 12),
    ("label", "Güven", 12),
    ("fake_breakout_risk", "Fake Breakout Riski", 16),
    ("last_price", "Güncel Fiyat", 14),
    ("breakout_level", "Breakout Seviyesi", 16),
    ("target", "Hedef", 14),
    ("stop_loss", "Stop Loss", 14),
]


def _write_sheet(wb: Workbook, title: str, rows: List[Dict]):
    ws = wb.create_sheet(title=title[:31])  # Excel sekme adı limiti 31 karakter

    # Başlık satırı
    for col_idx, (key, header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"

    # Veri satırları
    for row_idx, r in enumerate(rows, start=2):
        direction = r.get("direction", "")
        row_fill = LONG_FILL if direction == "LONG" else SHORT_FILL
        for col_idx, (key, header, width) in enumerate(COLUMNS, start=1):
            value = r.get(key)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.fill = row_fill
            cell.border = THIN_BORDER
            if key in ("score", "success_probability", "last_price",
                       "breakout_level", "target", "stop_loss"):
                cell.alignment = Alignment(horizontal="right")
                if key in ("last_price", "breakout_level", "target", "stop_loss") and value is not None:
                    cell.number_format = "#,##0.0000"
                elif key == "success_probability" and value is not None:
                    cell.number_format = '0.0"%"'
            else:
                cell.alignment = Alignment(horizontal="center")

    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    # Boş tablo durumunda bilgi notu
    if not rows:
        ws.cell(row=2, column=1, value="Bu kategoride sonuç bulunamadı.").font = Font(
            name=FONT_NAME, italic=True, size=10)


def build_excel(all_results: List[Dict], top_long: List[Dict], top_short: List[Dict]) -> io.BytesIO:
    """Sonuçlardan bir .xlsx dosyası oluşturur ve bellek üzerindeki (BytesIO) halini döndürür."""
    wb = Workbook()
    wb.remove(wb.active)  # varsayılan boş sekmeyi kaldır

    _write_sheet(wb, "Tüm Sonuçlar", all_results)
    _write_sheet(wb, "En Güçlü LONG", top_long)
    _write_sheet(wb, "En Güçlü SHORT", top_short)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def export_filename(version_label: str = "") -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    prefix = f"{version_label}_" if version_label else ""
    return f"{prefix}kripto_formasyon_taramasi_{ts}.xlsx"


# --------------------------------------------------------------------------
# Sinyal Günlüğü (journal) dışa aktarımı
# --------------------------------------------------------------------------

JOURNAL_COLUMNS = [
    ("created_at", "Oluşturulma Tarihi", 18),
    ("symbol", "Coin", 12),
    ("timeframe", "Zaman Dilimi", 12),
    ("pattern", "Formasyon", 26),
    ("direction", "Yön", 8),
    ("entry_price", "Giriş Fiyatı", 14),
    ("target", "Hedef", 14),
    ("stop_loss", "Stop Loss", 14),
    ("score", "Skor", 8),
    ("success_probability", "Tahmini Başarı %", 14),
    ("label", "Güven", 12),
    ("status", "Durum", 12),
    ("last_price", "Son Kontrol Fiyatı", 16),
    ("closed_at", "Kapanış Tarihi", 18),
    ("close_price", "Kapanış Fiyatı", 14),
]

STATUS_FILLS = {
    "AÇIK": PatternFill(start_color="FFF6D9", end_color="FFF6D9", fill_type="solid"),
    "HEDEF": PatternFill(start_color="E6F7F5", end_color="E6F7F5", fill_type="solid"),
    "STOP": PatternFill(start_color="FDECEC", end_color="FDECEC", fill_type="solid"),
}


def build_journal_excel(signals: List[Dict], summary: Dict, version_label: str = "") -> io.BytesIO:
    """Sinyal günlüğünü (journal.py) biçimlendirilmiş bir .xlsx dosyasına aktarır."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sinyal Günlüğü"

    # Üst bilgi: versiyon ve özet istatistikler
    ws.cell(row=1, column=1, value=f"Versiyon: {version_label or '-'}").font = Font(
        name=FONT_NAME, bold=True, size=12)
    ws.cell(row=2, column=1, value=(
        f"Toplam: {summary.get('toplam', 0)}  ·  Açık: {summary.get('açık', 0)}  ·  "
        f"Hedef: {summary.get('hedef', 0)}  ·  Stop: {summary.get('stop', 0)}  ·  "
        f"Gerçek Başarı Oranı: {summary.get('gerçek_başarı_oranı')}%"
    )).font = Font(name=FONT_NAME, size=10, italic=True)

    header_row = 4
    for col_idx, (key, header, width) in enumerate(JOURNAL_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = f"A{header_row + 1}"

    for row_offset, sig in enumerate(signals):
        row_idx = header_row + 1 + row_offset
        fill = STATUS_FILLS.get(sig.get("status"), None)
        for col_idx, (key, header, width) in enumerate(JOURNAL_COLUMNS, start=1):
            value = sig.get(key)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.border = THIN_BORDER
            if fill:
                cell.fill = fill
            if key in ("entry_price", "target", "stop_loss", "last_price", "close_price") and value is not None:
                cell.number_format = "#,##0.0000"
                cell.alignment = Alignment(horizontal="right")
            elif key == "success_probability" and value is not None:
                cell.number_format = '0.0"%"'
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="center")

    if signals:
        ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(JOURNAL_COLUMNS))}{header_row + len(signals)}"
        )

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def journal_export_filename(version_label: str = "") -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    prefix = f"{version_label}_" if version_label else ""
    return f"{prefix}sinyal_gunlugu_{ts}.xlsx"
